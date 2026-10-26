import shutil
from pathlib import Path

import docx
import pytest
from PIL import Image, ImageDraw, ImageFont

import text_extract as te

FIXTURE = Path(__file__).parent / "fixtures" / "contract_sample.pdf"


def test_pdf_text_layer():
    r = te.extract(FIXTURE)
    assert r.method == "pdf-text" and r.pages == 1
    assert "ООО \"Альфа\"" in r.text


def test_docx_includes_tables(tmp_path):
    d = docx.Document()
    d.add_paragraph("Договор №5")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "ИНН"
    t.rows[0].cells[1].text = "7701234567"
    p = tmp_path / "a.docx"
    d.save(p)
    r = te.extract(p)
    assert r.method == "docx" and "7701234567" in r.text


def test_normalize_hyphenation_and_spaces():
    assert te.normalize_text("основ-\nные   научные\u00a0результаты") == "основные научные результаты"


def test_unsupported_extension(tmp_path):
    p = tmp_path / "a.xyz"
    p.write_text("x")
    assert te.extract(p).method == "unsupported"


@pytest.mark.skipif(not shutil.which("tesseract"), reason="tesseract не установлен")
def test_scanned_pdf_triggers_ocr(tmp_path, monkeypatch):
    monkeypatch.setattr(te, "OCR_LANG", "eng")
    img = Image.new("RGB", (1200, 400), "white")
    ImageDraw.Draw(img).text((40, 150), "CONTRACT NUMBER 42", fill="black",
                             font=ImageFont.load_default(size=60))
    p = tmp_path / "scan.pdf"
    img.save(p, "PDF")  # PDF из одной картинки — текстового слоя нет
    r = te.extract(p)
    assert r.method == "pdf-ocr"
    assert "CONTRACT" in r.text.upper()
