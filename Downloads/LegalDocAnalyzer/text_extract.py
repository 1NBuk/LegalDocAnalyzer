"""Извлечение текста из PDF / DOCX / изображений.

Основные изменения относительно первой версии:
* PDF читается через PyMuPDF (в десятки раз быстрее pdfplumber, не задваивает буквы);
* если у PDF нет текстового слоя (скан) — автоматический OCR страниц;
* DOCX читается вместе с таблицами (в договорах там часто реквизиты);
* текст нормализуется: переносы слов, неразрывные пробелы, лишние пробелы.
"""
from __future__ import annotations

import logging
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

OCR_LANG = os.environ.get("OCR_LANG", "rus+eng")
OCR_MAX_PAGES = int(os.environ.get("OCR_MAX_PAGES", "30"))
OCR_DPI = int(os.environ.get("OCR_DPI", "200"))
# если в среднем на странице меньше этого числа символов — считаем PDF сканом
MIN_CHARS_PER_PAGE = 25

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}


@dataclass
class ExtractResult:
    text: str
    pages: int = 0
    method: str = "none"  # pdf-text | pdf-ocr | docx | image-ocr | unsupported | error
    error: Optional[str] = None


# --------------------------------------------------------------------------- #
# нормализация
# --------------------------------------------------------------------------- #
_HYPHEN_BREAK = re.compile(r"(?<=[A-Za-zА-Яа-яЁё])[-\u00ad]\n(?=[a-zа-яё])")
_SPACES = re.compile(r"[ \t\u00a0\u2009\u202f]+")
_BLANKS = re.compile(r"\n{3,}")


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = _HYPHEN_BREAK.sub("", text)  # «основ-\nные» -> «основные»
    text = _SPACES.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _BLANKS.sub("\n\n", text).strip()


# --------------------------------------------------------------------------- #
# OCR
# --------------------------------------------------------------------------- #
def _ocr_lang() -> str:
    """Подбирает доступные в tesseract языки из запрошенных (rus+eng)."""
    try:
        import pytesseract

        available = set(pytesseract.get_languages(config=""))
    except Exception:  # tesseract не установлен
        return OCR_LANG
    wanted = [x for x in OCR_LANG.split("+") if x in available]
    if wanted:
        return "+".join(wanted)
    fallback = [x for x in ("eng",) if x in available] or sorted(available - {"osd"})[:1]
    if fallback:
        logger.warning("OCR: языки %s недоступны, использую %s", OCR_LANG, fallback[0])
        return fallback[0]
    return OCR_LANG


def _ocr_image(img) -> str:
    import pytesseract

    return pytesseract.image_to_string(img, lang=_ocr_lang())


def _ocr_pdf(doc) -> str:
    from PIL import Image

    parts = []
    for i, page in enumerate(doc):
        if i >= OCR_MAX_PAGES:
            logger.warning("OCR: обработаны первые %d страниц из %d", OCR_MAX_PAGES, len(doc))
            break
        pix = page.get_pixmap(dpi=OCR_DPI)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        parts.append(_ocr_image(img))
    return "\n".join(parts)


# --------------------------------------------------------------------------- #
# форматы
# --------------------------------------------------------------------------- #
def extract_text_from_pdf(path: Path) -> ExtractResult:
    import pymupdf

    try:
        with pymupdf.open(path) as doc:
            pages = len(doc)
            text = "\n".join(page.get_text("text") for page in doc)
            if pages and len(text.strip()) / pages < MIN_CHARS_PER_PAGE:
                logger.info("PDF %s без текстового слоя — запускаю OCR", path.name)
                try:
                    return ExtractResult(normalize_text(_ocr_pdf(doc)), pages, "pdf-ocr")
                except Exception as e:  # tesseract отсутствует и т.п.
                    logger.error("OCR не удался для %s: %s", path.name, e)
                    return ExtractResult(normalize_text(text), pages, "pdf-ocr", f"ocr_failed: {e}")
            return ExtractResult(normalize_text(text), pages, "pdf-text")
    except Exception as e:
        logger.exception("PDF extract error %s", path)
        return ExtractResult("", 0, "error", f"pdf_error: {e}")


def extract_text_from_docx(path: Path) -> ExtractResult:
    import docx

    try:
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    if cell.text not in cells:  # объединённые ячейки дублируются
                        cells.append(cell.text)
                parts.append(" | ".join(cells))
        return ExtractResult(normalize_text("\n".join(parts)), 1, "docx")
    except Exception as e:
        logger.exception("DOCX extract error %s", path)
        return ExtractResult("", 0, "error", f"docx_error: {e}")


def extract_text_from_image(path: Path) -> ExtractResult:
    from PIL import Image

    try:
        with Image.open(path) as img:
            return ExtractResult(normalize_text(_ocr_image(img)), 1, "image-ocr")
    except Exception as e:
        logger.exception("OCR error %s", path)
        return ExtractResult("", 0, "error", f"ocr_error: {e}")


def extract(path: Path) -> ExtractResult:
    """Главная точка входа: текст + метаданные (страницы, способ, ошибка)."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(path)
    if ext == ".docx":
        return extract_text_from_docx(path)
    if ext in IMAGE_EXT:
        return extract_text_from_image(path)
    return ExtractResult("", 0, "unsupported", f"unsupported_extension: {ext}")


def extract_text(path: Path) -> str:
    """Обратно-совместимый вариант: только строка текста."""
    return extract(path).text
